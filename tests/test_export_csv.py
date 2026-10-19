"""Tests for the in-memory CSV ZIP export."""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import json
import zipfile

import openpyxl
import pandas as pd
import pytest

from src.app.export_csv import render_analysis_csv_zip


@pytest.fixture
def export_context():
    geometry = {
        "type": "Polygon",
        "coordinates": [[[-52.123456789, -28.0], [-51.9, -28.0], [-51.9, -27.9], [-52.123456789, -28.0]]],
        "crs": {"type": "name", "properties": {"name": "EPSG:4326"}},
    }
    return {
        "period": {"start_date": "2025-01-01", "end_date": "2025-03-31"},
        "area_ha": 3.2,
        "geometry_geojson": geometry,
        "map_center": [-28.0, -51.95],
        "image_count": 12,
        "climate_summary": {
            "total_precipitation_mm": 82.5,
            "mean_temperature_c": 23.1,
            "mean_solar_radiation_mj_m2_day": 18.7,
        },
        "climate_data": pd.DataFrame([
            {"date": "2025-01-10", "precipitation": 12.5, "temperature": 22.0, "solar_radiation": 4.5},
            {"date": "2025-02-10", "precipitation": None, "temperature": 24.0, "solar_radiation": None},
        ]),
        "indices": {
            name: {
                "trend": "stable",
                "time_series": [
                    {"date": "2025-01-10", "value": 0.42},
                    {"date": "2025-02-10", "value": None},
                ],
                "anomalies": [
                    {"date": "2025-01-10", "value": 0.42, "anomaly": True, "classification": "low"},
                    {"date": "2025-02-10", "value": None, "anomaly": False, "classification": None},
                ],
            }
            for name in ("NDVI", "NDWI", "NDMI")
        },
    }


def read_zip(csv_zip):
    archive = zipfile.ZipFile(BytesIO(csv_zip))
    return archive, {
        name: list(csv.DictReader(StringIO(archive.read(name).decode("utf-8"))))
        for name in archive.namelist()
    }


def test_zip_contains_only_the_seven_expected_csvs(export_context):
    archive, _ = read_zip(render_analysis_csv_zip(export_context))
    assert archive.namelist() == [
        "resumo.csv", "ndvi.csv", "ndwi.csv", "ndmi.csv", "clima.csv",
        "anomalias.csv", "geografia.csv",
    ]
    assert all("/" not in name and "\\" not in name for name in archive.namelist())


def test_csv_headers_are_stable_and_data_matches_contract(export_context):
    _, data = read_zip(render_analysis_csv_zip(export_context, generated_at=date(2025, 4, 1)))
    assert list(data["resumo.csv"][0]) == ["campo", "valor", "unidade"]
    for name in ("ndvi.csv", "ndwi.csv", "ndmi.csv"):
        assert list(data[name][0]) == ["date", "value", "trend", "anomaly"]
    assert list(data["clima.csv"][0]) == [
        "date", "precipitation_mm", "temperature_c", "solar_radiation_mj_m2_day"
    ]
    assert list(data["anomalias.csv"][0]) == [
        "index", "date", "value", "anomaly", "classification"
    ]
    assert list(data["geografia.csv"][0]) == ["field", "value", "unit"]
    summary = {r["campo"]: (r["valor"], r["unidade"]) for r in data["resumo.csv"]}
    assert summary["periodo_inicial"] == ("2025-01-01", "")
    assert summary["area_ha"] == ("3.2", "ha")
    assert summary["data_geracao"] == ("2025-04-01", "")


def test_index_climate_anomalies_and_geography_are_serialized(export_context):
    _, data = read_zip(render_analysis_csv_zip(export_context))
    assert data["ndvi.csv"][0] == {
        "date": "2025-01-10", "value": "0.42", "trend": "stable", "anomaly": "true"
    }
    assert data["ndvi.csv"][1]["value"] == ""
    assert data["ndvi.csv"][1]["anomaly"] == "false"
    assert data["clima.csv"][0]["precipitation_mm"] == "12.5"
    assert data["clima.csv"][1]["precipitation_mm"] == ""
    assert len(data["anomalias.csv"]) == 3
    geography = {r["field"]: r["value"] for r in data["geografia.csv"]}
    assert geography["crs"] == "EPSG:4326"
    assert json.loads(geography["geojson"]) == export_context["geometry_geojson"]
    assert "-52.123456789" in geography["geojson"]


def test_empty_sets_keep_headers_and_zip_is_valid(export_context):
    export_context["climate_data"] = pd.DataFrame()
    for result in export_context["indices"].values():
        result["time_series"] = []
        result["anomalies"] = []
    archive, data = read_zip(render_analysis_csv_zip(export_context))
    assert archive.testzip() is None
    assert data["clima.csv"] == []
    assert data["ndvi.csv"] == []
    assert data["anomalias.csv"] == []
    assert archive.read("clima.csv").decode("utf-8").splitlines() == [
        "date,precipitation_mm,temperature_c,solar_radiation_mj_m2_day"
    ]


@pytest.mark.parametrize("indices", [None, {}, {"NDVI": {}, "NDWI": {}}])
def test_invalid_or_incomplete_context_is_rejected(export_context, indices):
    export_context["indices"] = indices
    with pytest.raises(ValueError, match="NDVI, NDWI e NDMI"):
        render_analysis_csv_zip(export_context)


def test_non_dictionary_context_is_rejected():
    with pytest.raises(ValueError, match="dicionario"):
        render_analysis_csv_zip(None)


def test_generation_is_in_memory_utf8_and_does_not_make_network_calls(export_context, monkeypatch):
    import requests

    monkeypatch.setattr(requests, "get", lambda *a, **k: pytest.fail("unexpected HTTP request"))
    payload = render_analysis_csv_zip(export_context)
    assert isinstance(payload, bytes)
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        assert all(info.filename == info.filename.lstrip("/") for info in archive.infolist())
        assert all(info.file_size > 0 for info in archive.infolist())


def test_csv_and_xlsx_have_equivalent_climate_and_summary_values(export_context):
    from src.app.export_xlsx import render_analysis_xlsx

    generated = datetime(2025, 4, 1)
    _, csv_data = read_zip(render_analysis_csv_zip(export_context, generated_at=generated))
    workbook = openpyxl.load_workbook(BytesIO(render_analysis_xlsx(export_context, generated_at=generated)))

    csv_climate = csv_data["clima.csv"]
    xlsx_climate = list(workbook["Clima"].iter_rows(min_row=2, values_only=True))
    assert [row["date"] for row in csv_climate] == [row[0] for row in xlsx_climate]
    for csv_row, xlsx_row in zip(csv_climate, xlsx_climate):
        assert [float(csv_row[key]) if csv_row[key] else None for key in (
            "precipitation_mm", "temperature_c", "solar_radiation_mj_m2_day"
        )] == list(xlsx_row[1:])

    csv_summary = {row["campo"]: row["valor"] for row in csv_data["resumo.csv"]}
    xlsx_summary = {
        row[0]: row[1]
        for row in workbook["Resumo"].iter_rows(min_row=2, values_only=True)
    }
    assert csv_summary["precipitacao_total_mm"] == str(xlsx_summary["precipitacao_total_mm"])
    assert csv_summary["temperatura_media_c"] == str(xlsx_summary["temperatura_media_c"])
    assert csv_summary["radiacao_solar_media_mj_m2_dia"] == str(xlsx_summary["radiacao_solar_media_mj_m2_dia"])


def test_csv_scalar_and_record_normalizers_cover_null_numeric_and_container_cases():
    from datetime import datetime
    from src.app.export_csv import _cell, _number, _records

    assert _number("2.5") == 2.5
    assert _number("bad") is None
    assert _number(float("inf")) is None
    assert _number(True) is None
    assert _cell(True) == "true"
    assert _cell(False) == "false"
    assert _cell(datetime(2025, 1, 2, 3)) == "2025-01-02"
    assert _cell(float("nan")) == ""
    assert _cell(pd.NA) == ""
    assert _cell([1, 2]) == "[1, 2]"
    assert _records(None) == []
    assert _records({"x": 1}) == [{"x": 1}]
    assert _records(({"x": 1}, "ignored")) == [{"x": 1}]
    assert _records("unsupported") == []


def test_csv_climate_accepts_normalized_fields_and_missing_radiation(export_context):
    export_context["climate_data"] = [
        {"date": "2025-01-10", "precipitation_mm": "3.5", "temperature_c": "21", "solar_radiation_mj_m2_day": "18"},
        {"date": "2025-01-11", "precipitation": None, "temperature": None, "solar_radiation": "invalid"},
    ]
    _, data = read_zip(render_analysis_csv_zip(export_context))
    assert data["clima.csv"][0]["solar_radiation_mj_m2_day"] == "18.0"
    assert data["clima.csv"][1]["solar_radiation_mj_m2_day"] == ""


def test_csv_skips_anomaly_with_missing_date_and_missing_map_center(export_context):
    export_context["map_center"] = [1]
    export_context["indices"]["NDVI"]["anomalies"] = [
        {"anomaly": True, "date": None, "value": 0.5},
    ]
    export_context["indices"]["NDWI"]["anomalies"] = []
    export_context["indices"]["NDMI"]["anomalies"] = []
    _, data = read_zip(render_analysis_csv_zip(export_context))
    assert data["ndvi.csv"][0]["anomaly"] == "false"
    assert len(data["anomalias.csv"]) == 1
    assert data["anomalias.csv"][0]["date"] == ""
    geo_fields = [row["field"] for row in data["geografia.csv"]]
    assert "centroid_latitude" not in geo_fields


def test_xlsx_handles_true_anomaly_without_date_and_invalid_radiation(export_context):
    from src.app.export_xlsx import render_analysis_xlsx

    export_context["indices"]["NDVI"]["anomalies"] = [
        {"anomaly": True, "date": None, "value": None},
    ]
    export_context["indices"]["NDWI"]["anomalies"] = []
    export_context["indices"]["NDMI"]["anomalies"] = []
    export_context["climate_data"] = pd.DataFrame([
        {"date": "2025-01-10", "solar_radiation": "invalid"},
    ])
    workbook = openpyxl.load_workbook(BytesIO(render_analysis_xlsx(export_context)))
    assert workbook["NDVI"]["D2"].value is False
    assert workbook["Anomalias"].max_row == 2
    assert workbook["Clima"]["D2"].value is None
