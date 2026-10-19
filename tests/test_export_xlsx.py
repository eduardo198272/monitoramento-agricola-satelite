"""Tests for XLSX export builder."""

from datetime import date
from io import BytesIO

import openpyxl
import pytest

from src.app.export_xlsx import render_analysis_xlsx


@pytest.fixture
def export_context():
    """A complete export context with all three indices and typical data."""
    return {
        "period": {"start_date": "2025-01-01", "end_date": "2025-03-31"},
        "area_ha": 3.2,
        "geometry_geojson": {
            "type": "Polygon",
            "coordinates": [[[-52, -28], [-51.9, -28], [-51.9, -27.9], [-52, -28]]],
            "crs": {"type": "name", "properties": {"name": "EPSG:4326"}},
        },
        "map_center": [-28.0, -51.95],
        "image_count": 12,
        "climate_summary": {
            "precipitation_mm": 82.5,
            "temperature_c": 23.1,
            "solar_radiation_mj_m2_day": 18.7,
        },
        "climate_data": [
            {"date": "2025-01-10", "precipitation_mm": 12.5, "temperature_c": 22.0, "solar_radiation_mj_m2_day": 16.2},
            {"date": "2025-02-10", "precipitation_mm": 35.0, "temperature_c": 24.0, "solar_radiation_mj_m2_day": 19.1},
            {"date": "2025-03-10", "precipitation_mm": 35.0, "temperature_c": 23.5, "solar_radiation_mj_m2_day": 21.0},
        ],
        "indices": {
            "NDVI": {
                "mean_value": 0.42,
                "trend": "stable",
                "time_series": [
                    {"date": "2025-01-10", "value": 0.38},
                    {"date": "2025-02-10", "value": 0.42},
                    {"date": "2025-03-10", "value": 0.46},
                ],
                "anomalies": [
                    {"date": "2025-02-10", "value": 0.42, "anomaly": True, "classification": "high", "z_score": 1.5},
                ],
                "index_map": object(),
            },
            "NDWI": {
                "mean_value": 0.12,
                "trend": "increasing",
                "time_series": [
                    {"date": "2025-01-10", "value": 0.08},
                    {"date": "2025-02-10", "value": 0.12},
                    {"date": "2025-03-10", "value": 0.16},
                ],
                "anomalies": [],
                "index_map": object(),
            },
            "NDMI": {
                "mean_value": 0.28,
                "trend": "decreasing",
                "time_series": [
                    {"date": "2025-01-10", "value": 0.32},
                    {"date": "2025-02-10", "value": 0.28},
                    {"date": "2025-03-10", "value": 0.24},
                ],
                "anomalies": [
                    {"date": "2025-01-10", "value": 0.32, "anomaly": True, "classification": "low"},
                    {"date": "2025-03-10", "value": 0.24, "anomaly": True, "classification": "low"},
                ],
                "index_map": object(),
            },
        },
    }


def test_render_analysis_xlsx_returns_valid_bytes(export_context):
    """XLSX renderer returns valid workbook bytes."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    assert isinstance(xlsx_bytes, bytes)
    assert xlsx_bytes.startswith(b"PK")  # ZIP signature


def test_render_analysis_xlsx_has_expected_sheets(export_context):
    """Workbook contains exactly the 7 expected sheets."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    expected_sheets = ["Resumo", "NDVI", "NDWI", "NDMI", "Clima", "Anomalias", "Geografia"]
    assert wb.sheetnames == expected_sheets


def test_resumo_sheet_structure_and_content(export_context):
    """Resumo sheet has campo, valor, unidade columns with expected rows."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Resumo"]

    # Header row
    assert [cell.value for cell in ws[1]] == ["campo", "valor", "unidade"]

    # Collect all rows
    rows = {row[0].value: (row[1].value, row[2].value) for row in ws.iter_rows(min_row=2, values_only=False) if row[0].value}

    # Required fields
    assert "periodo_inicial" in rows
    assert rows["periodo_inicial"] == ("2025-01-01", None)
    assert "periodo_final" in rows
    assert rows["periodo_final"] == ("2025-03-31", None)
    assert "area_ha" in rows
    assert rows["area_ha"] == (3.2, "ha")
    assert "imagens_analisadas" in rows
    assert rows["imagens_analisadas"] == (12, "imagens")
    assert "data_geracao" in rows  # date when generated
    assert "precipitacao_total_mm" in rows
    assert rows["precipitacao_total_mm"] == (82.5, "mm")
    assert "temperatura_media_c" in rows
    assert rows["temperatura_media_c"] == (23.1, "C")
    assert "radiacao_solar_media_mj_m2_dia" in rows
    assert rows["radiacao_solar_media_mj_m2_dia"] == (18.7, "MJ/m2/dia")


def test_index_sheets_structure_and_content(export_context):
    """Each index sheet (NDVI, NDWI, NDMI) has date, value, trend, anomaly columns."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    for index_name in ("NDVI", "NDWI", "NDMI"):
        ws = wb[index_name]
        # Header
        assert [cell.value for cell in ws[1]] == ["date", "value", "trend", "anomaly"]

        # Data rows (3 per index)
        data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        assert len(data_rows) == 3

        # Check first row structure
        for row in data_rows:
            assert isinstance(row[0], str)  # date as string
            assert row[1] is not None and isinstance(row[1], (int, float))  # value
            assert row[2] in ("stable", "increasing", "decreasing")  # trend
            assert row[3] in (True, False, None)  # anomaly flag


def test_index_sheet_trend_and_anomaly_per_row(export_context):
    """Trend and anomaly info is correctly associated with each date row."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    # NDVI: trend="stable", one anomaly on 2025-02-10
    ws_ndvi = wb["NDVI"]
    ndvi_rows = {row[0]: row for row in ws_ndvi.iter_rows(min_row=2, values_only=True) if row[0]}
    assert ndvi_rows["2025-01-10"][2] == "stable"
    assert ndvi_rows["2025-01-10"][3] is False
    assert ndvi_rows["2025-02-10"][2] == "stable"
    assert ndvi_rows["2025-02-10"][3] is True
    assert ndvi_rows["2025-03-10"][2] == "stable"
    assert ndvi_rows["2025-03-10"][3] is False


def test_clima_sheet_structure_and_content(export_context):
    """Clima sheet has date, precipitation_mm, temperature_c, solar_radiation_mj_m2_day."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Clima"]

    # Header
    assert [cell.value for cell in ws[1]] == [
        "date", "precipitation_mm", "temperature_c", "solar_radiation_mj_m2_day"
    ]

    # Data rows
    data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 3

    for row in data_rows:
        assert isinstance(row[0], str)  # date
        # Values can be float or None
        assert row[1] is None or isinstance(row[1], (int, float))
        assert row[2] is None or isinstance(row[2], (int, float))
        assert row[3] is None or isinstance(row[3], (int, float))


def test_clima_sheet_preserves_nulls(export_context):
    """Clima sheet keeps None for missing values."""
    export_context["climate_data"] = [
        {"date": "2025-01-10", "precipitation_mm": None, "temperature_c": 22.0, "solar_radiation_mj_m2_day": 16.2},
        {"date": "2025-02-10", "precipitation_mm": 35.0, "temperature_c": None, "solar_radiation_mj_m2_day": None},
    ]
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Clima"]

    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert rows[0][1] is None  # precipitation None
    assert rows[1][2] is None  # temperature None
    assert rows[1][3] is None  # radiation None


def test_anomalias_sheet_structure_and_content(export_context):
    """Anomalias sheet has index, date, value, anomaly, classification columns."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Anomalias"]

    # Header
    assert [cell.value for cell in ws[1]] == ["index", "date", "value", "anomaly", "classification"]

    # Data rows: NDVI has 1, NDMI has 2 = 3 total
    data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 3

    # Check NDVI anomaly
    ndvi_anomaly = next(r for r in data_rows if r[0] == "NDVI")
    assert ndvi_anomaly[1] == "2025-02-10"
    assert ndvi_anomaly[2] == 0.42
    assert ndvi_anomaly[3] is True
    assert ndvi_anomaly[4] == "high"

    # Check NDMI anomalies
    ndmi_anomalies = [r for r in data_rows if r[0] == "NDMI"]
    assert len(ndmi_anomalies) == 2
    assert ndmi_anomalies[0][4] == "low"
    assert ndmi_anomalies[1][4] == "low"


def test_geografia_sheet_structure_and_content(export_context):
    """Geografia sheet has field, value, unit with area, centroid, CRS, GeoJSON."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Geografia"]

    # Header
    assert [cell.value for cell in ws[1]] == ["field", "value", "unit"]

    rows = {row[0]: (row[1], row[2]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}

    assert rows["area_ha"] == (3.2, "ha")
    assert rows["centroid_latitude"] == (-28.0, "graus")
    assert rows["centroid_longitude"] == (-51.95, "graus")
    assert rows["crs"] == ("EPSG:4326", None)
    assert "geojson" in rows
    geojson_value = rows["geojson"][0]
    assert isinstance(geojson_value, str)
    assert "Polygon" in geojson_value
    assert "[-52,-28]" in geojson_value


def test_xlsx_preserves_numeric_types_not_strings(export_context):
    """Numeric values are stored as numbers, not strings, in XLSX cells."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    # Check Resumo area_ha is numeric
    ws_resumo = wb["Resumo"]
    for row in ws_resumo.iter_rows(min_row=2, values_only=True):
        if row[0] == "area_ha":
            assert isinstance(row[1], (int, float))
            break

    # Check index values are numeric
    for idx in ("NDVI", "NDWI", "NDMI"):
        ws = wb[idx]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                assert isinstance(row[1], (int, float))


def test_xlsx_dates_are_iso_format_strings(export_context):
    """Dates in XLSX are ISO format strings (YYYY-MM-DD)."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    for idx in ("NDVI", "NDWI", "NDMI"):
        ws = wb[idx]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                assert isinstance(row[0], str)
                assert len(row[0]) == 10  # YYYY-MM-DD
                assert row[0].count("-") == 2


def test_render_analysis_xlsx_rejects_missing_indices(export_context):
    """Renderer rejects context missing any of the three indices."""
    export_context["indices"].pop("NDMI")
    with pytest.raises(ValueError, match="NDVI, NDWI e NDMI"):
        render_analysis_xlsx(export_context)


def test_render_analysis_xlsx_rejects_non_dict_context():
    """Renderer rejects non-dict context."""
    with pytest.raises(ValueError, match="dicionario"):
        render_analysis_xlsx(None)


def test_render_analysis_xlsx_handles_empty_time_series(export_context):
    """XLSX handles empty time series gracefully (headers only)."""
    export_context["indices"]["NDVI"]["time_series"] = []
    export_context["indices"]["NDWI"]["time_series"] = []
    export_context["indices"]["NDMI"]["time_series"] = []

    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    for idx in ("NDVI", "NDWI", "NDMI"):
        ws = wb[idx]
        data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        assert len(data_rows) == 0  # Only header


def test_render_analysis_xlsx_handles_empty_climate_data(export_context):
    """XLSX handles empty climate data (headers only)."""
    export_context["climate_data"] = []

    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Clima"]

    data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 0


def test_render_analysis_xlsx_handles_empty_anomalies(export_context):
    """XLSX handles no anomalies (headers only)."""
    export_context["indices"]["NDVI"]["anomalies"] = []
    export_context["indices"]["NDMI"]["anomalies"] = []

    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Anomalias"]

    data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 0


def test_render_analysis_xlsx_handles_missing_optional_fields(export_context):
    """XLSX handles missing optional fields like area_ha, climate_summary."""
    export_context["area_ha"] = None
    export_context["climate_summary"] = {}
    export_context["map_center"] = None
    export_context["image_count"] = None

    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    ws_resumo = wb["Resumo"]
    rows = {row[0].value: row[1].value for row in ws_resumo.iter_rows(min_row=2, values_only=False) if row[0].value}
    assert rows.get("area_ha") is None
    assert rows.get("imagens_analisadas") is None
    assert rows.get("precipitacao_total_mm") is None
    assert rows.get("temperatura_media_c") is None
    assert rows.get("radiacao_solar_media_mj_m2_dia") is None

    ws_geografia = wb["Geografia"]
    geo_rows = {row[0]: row[1] for row in ws_geografia.iter_rows(min_row=2, values_only=True) if row[0]}
    assert geo_rows.get("centroid_latitude") is None
    assert geo_rows.get("centroid_longitude") is None


def test_xlsx_workbook_properties(export_context):
    """Workbook has proper metadata."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    # Check workbook properties
    assert wb.properties.title == "Relatório de Monitoramento Agrícola"
    assert wb.properties.creator == "Monitoramento Agrícola por Satélite"


def test_xlsx_sheet_order_is_fixed(export_context):
    """Sheet order is fixed: Resumo, NDVI, NDWI, NDMI, Clima, Anomalias, Geografia."""
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Resumo", "NDVI", "NDWI", "NDMI", "Clima", "Anomalias", "Geografia"]


def test_xlsx_no_external_calls_during_generation(export_context, monkeypatch):
    """XLSX generation performs no HTTP requests."""
    import requests
    monkeypatch.setattr(requests, "get", lambda *a, **k: pytest.fail("HTTP request during XLSX"))
    xlsx_bytes = render_analysis_xlsx(export_context)
    assert xlsx_bytes.startswith(b"PK")


def test_format_number_for_cell_handles_boolean(export_context):
    """_format_number_for_cell returns None for boolean values."""
    from src.app.export_xlsx import _format_number_for_cell
    assert _format_number_for_cell(True) is None
    assert _format_number_for_cell(False) is None


def test_format_number_for_cell_handles_string_number(export_context):
    """_format_number_for_cell converts numeric strings."""
    from src.app.export_xlsx import _format_number_for_cell
    assert _format_number_for_cell("0.42") == 0.42
    assert _format_number_for_cell("123") == 123.0
    assert _format_number_for_cell("not-a-number") is None


def test_to_iso_date_handles_various_formats(export_context):
    """_to_iso_date handles various date formats."""
    from src.app.export_xlsx import _to_iso_date
    from datetime import date, datetime

    # None
    assert _to_iso_date(None) is None
    # ISO string
    assert _to_iso_date("2025-01-10") == "2025-01-10"
    # YYYYMMDD format
    assert _to_iso_date("20250110") == "2025-01-10"
    # Invalid YYYYMMDD returns original
    assert _to_iso_date("20259999") == "20259999"
    # date object
    assert _to_iso_date(date(2025, 1, 10)) == "2025-01-10"
    # datetime object - returns full ISO with time
    assert _to_iso_date(datetime(2025, 1, 10, 12, 0)) == "2025-01-10T12:00:00"
    # Object with isoformat method
    class WithIsoformat:
        def isoformat(self):
            return "2025-01-10T12:00:00"
    assert _to_iso_date(WithIsoformat()) == "2025-01-10T12:00:00"
    # Fallback to str
    assert _to_iso_date(12345) == "12345"


def test_index_sheet_handles_anomaly_with_true_flag(export_context):
    """Index sheet correctly marks anomaly flag when anomaly.anomaly is True."""
    export_context["indices"]["NDVI"]["time_series"] = [
        {"date": "2025-01-10", "value": 0.38},
        {"date": "2025-02-10", "value": 0.42},
    ]
    export_context["indices"]["NDVI"]["anomalies"] = [
        {"date": "2025-02-10", "value": 0.42, "anomaly": True, "classification": "high"},
        {"date": "2025-01-10", "value": 0.38, "anomaly": False, "classification": "none"},
    ]
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["NDVI"]
    rows = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert rows["2025-01-10"][3] is False
    assert rows["2025-02-10"][3] is True


def test_index_sheet_handles_anomaly_with_invalid_date(export_context):
    """Index sheet handles anomaly with invalid/unparsable date gracefully."""
    export_context["indices"]["NDVI"]["anomalies"] = [
        {"date": "invalid-date", "value": 0.42, "anomaly": True, "classification": "high"},
    ]
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["NDVI"]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    # Anomaly with invalid date won't be in anomaly_by_date, so anomaly flag stays False
    for row in rows:
        assert row[3] is False


def test_index_sheet_handles_none_trend(export_context):
    """Index sheet handles missing trend (None)."""
    export_context["indices"]["NDVI"]["trend"] = None
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["NDVI"]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    for row in rows:
        assert row[2] is None


def test_clima_sheet_handles_missing_keys(export_context):
    """Clima sheet handles records missing some keys."""
    export_context["climate_data"] = [
        {"date": "2025-01-10", "precipitation_mm": 10.0},  # missing temperature, radiation
        {"date": "2025-02-10", "temperature_c": 22.0},  # missing precipitation, radiation
    ]
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Clima"]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert len(rows) == 2
    assert rows[0][1] == 10.0  # precipitation
    assert rows[0][2] is None  # temperature missing
    assert rows[0][3] is None  # radiation missing
    assert rows[1][1] is None  # precipitation missing
    assert rows[1][2] == 22.0  # temperature
    assert rows[1][3] is None  # radiation missing


def test_anomalias_sheet_skips_non_true_anomaly(export_context):
    """Anomalias sheet only includes anomalies where anomaly is True."""
    export_context["indices"]["NDVI"]["anomalies"] = [
        {"date": "2025-01-10", "value": 0.38, "anomaly": False, "classification": "none"},
        {"date": "2025-02-10", "value": 0.42, "anomaly": True, "classification": "high"},
        {"date": "2025-03-10", "value": 0.46, "anomaly": None, "classification": "unknown"},
    ]
    export_context["indices"]["NDMI"]["anomalies"] = []  # Clear fixture anomalies
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Anomalias"]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert len(rows) == 1
    assert rows[0][0] == "NDVI"
    assert rows[0][1] == "2025-02-10"
    assert rows[0][3] is True


def test_geografia_sheet_handles_missing_map_center(export_context):
    """Geografia sheet handles missing map_center gracefully."""
    export_context["map_center"] = None
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Geografia"]
    rows = {row[0]: (row[1], row[2]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert "centroid_latitude" not in rows
    assert "centroid_longitude" not in rows
    assert rows["crs"][0] == "EPSG:4326"


def test_geografia_sheet_handles_geometry_without_crs(export_context):
    """Geografia sheet handles geometry without CRS."""
    export_context["geometry_geojson"] = {
        "type": "Polygon",
        "coordinates": [[[-52, -28], [-51.9, -28], [-51.9, -27.9], [-52, -28]]],
    }
    xlsx_bytes = render_analysis_xlsx(export_context)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb["Geografia"]
    rows = {row[0]: (row[1], row[2]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert rows["crs"][0] == "EPSG:4326"


def test_render_analysis_xlsx_rejects_empty_indices_dict(export_context):
    """Renderer rejects context with empty indices dict."""
    export_context["indices"] = {}
    with pytest.raises(ValueError, match="NDVI, NDWI e NDMI"):
        render_analysis_xlsx(export_context)


def test_render_analysis_xlsx_rejects_none_indices(export_context):
    """Renderer rejects context with None indices."""
    export_context["indices"] = None
    with pytest.raises(ValueError, match="NDVI, NDWI e NDMI"):
        render_analysis_xlsx(export_context)


def test_xlsx_generated_at_uses_default_when_none(export_context):
    """XLSX uses current datetime when generated_at is None."""
    # Just ensure it doesn't crash and produces valid output
    xlsx_bytes = render_analysis_xlsx(export_context, generated_at=None)
    assert xlsx_bytes.startswith(b"PK")
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Resumo", "NDVI", "NDWI", "NDMI", "Clima", "Anomalias", "Geografia"]