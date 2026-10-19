"""XLSX workbook builder for normalized agricultural analysis snapshots."""

from datetime import date, datetime
from io import BytesIO
import json
from typing import Any

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, numbers

from src.app.export_context import SUPPORTED_INDICES


def _format_number_for_cell(value: Any) -> float | int | None:
    """Convert a value to a numeric type for XLSX cell, or None if not numeric."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_iso_date(value: Any) -> str | None:
    """Normalize a date value to ISO format string YYYY-MM-DD."""
    if value is None:
        return None
    if isinstance(value, str):
        # Already in ISO format or YYYYMMDD
        if len(value) == 10 and value.count("-") == 2:
            return value
        if len(value) == 8 and value.isdigit():
            try:
                return datetime.strptime(value, "%Y%m%d").date().isoformat()
            except ValueError:
                return value
        return value
    if isinstance(value, date):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _write_header(ws, headers: list[str], row: int = 1) -> None:
    """Write header row with bold font."""
    bold_font = Font(bold=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = bold_font


def _write_resumo_sheet(ws, context: dict, generated_at: date | datetime | None = None) -> None:
    """Write the Resumo sheet with campo, valor, unidade columns."""
    _write_header(ws, ["campo", "valor", "unidade"])

    generated_at = generated_at or datetime.now()
    period = context.get("period", {})
    area_ha = context.get("area_ha")
    image_count = context.get("image_count")
    climate_summary = context.get("climate_summary", {}) or {}

    rows = [
        ("periodo_inicial", period.get("start_date", ""), ""),
        ("periodo_final", period.get("end_date", ""), ""),
        ("area_ha", area_ha, "ha" if area_ha is not None else ""),
        ("imagens_analisadas", image_count, "imagens" if image_count is not None else ""),
        ("data_geracao", generated_at.date().isoformat() if hasattr(generated_at, "date") else generated_at.isoformat(), ""),
        ("precipitacao_total_mm", climate_summary.get("precipitation_mm"), "mm"),
        ("temperatura_media_c", climate_summary.get("temperature_c"), "C"),
        ("radiacao_solar_media_mj_m2_dia", climate_summary.get("solar_radiation_mj_m2_day"), "MJ/m2/dia"),
    ]

    for row_idx, (campo, valor, unidade) in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=campo)
        cell_valor = ws.cell(row=row_idx, column=2, value=_format_number_for_cell(valor))
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            cell_valor.number_format = "0.00"
        # For string values (like dates), write directly
        if not isinstance(valor, (int, float)) or isinstance(valor, bool):
            ws.cell(row=row_idx, column=2, value=valor)
        ws.cell(row=row_idx, column=3, value=unidade or "")


def _write_index_sheet(ws, index_name: str, index_data: dict) -> None:
    """Write an index sheet (NDVI, NDWI, NDMI) with date, value, trend, anomaly."""
    _write_header(ws, ["date", "value", "trend", "anomaly"])

    time_series = index_data.get("time_series", []) or []
    trend = index_data.get("trend")
    anomalies = index_data.get("anomalies", []) or []

    # Build anomaly lookup by date
    anomaly_by_date = {}
    for anomaly in anomalies:
        if anomaly.get("anomaly") is True:
            anomaly_date = _to_iso_date(anomaly.get("date"))
            if anomaly_date:
                anomaly_by_date[anomaly_date] = True

    for row_idx, record in enumerate(time_series, 2):
        date_value = _to_iso_date(record.get("date"))
        value = _format_number_for_cell(record.get("value"))
        anomaly_flag = anomaly_by_date.get(date_value, False)

        ws.cell(row=row_idx, column=1, value=date_value)
        cell_valor = ws.cell(row=row_idx, column=2, value=value)
        if value is not None:
            cell_valor.number_format = "0.00000000"
        ws.cell(row=row_idx, column=3, value=trend)
        ws.cell(row=row_idx, column=4, value=anomaly_flag)


def _write_clima_sheet(ws, climate_data: list[dict]) -> None:
    """Write the Clima sheet with date, precipitation_mm, temperature_c, solar_radiation_mj_m2_day."""
    _write_header(ws, ["date", "precipitation_mm", "temperature_c", "solar_radiation_mj_m2_day"])

    for row_idx, record in enumerate(climate_data or [], 2):
        date_value = _to_iso_date(record.get("date"))
        precip = _format_number_for_cell(record.get("precipitation_mm"))
        temp = _format_number_for_cell(record.get("temperature_c"))
        rad = _format_number_for_cell(record.get("solar_radiation_mj_m2_day"))

        ws.cell(row=row_idx, column=1, value=date_value)
        cell_precip = ws.cell(row=row_idx, column=2, value=precip)
        if precip is not None:
            cell_precip.number_format = "0.00"
        cell_temp = ws.cell(row=row_idx, column=3, value=temp)
        if temp is not None:
            cell_temp.number_format = "0.00"
        cell_rad = ws.cell(row=row_idx, column=4, value=rad)
        if rad is not None:
            cell_rad.number_format = "0.00"


def _write_anomalias_sheet(ws, indices: dict) -> None:
    """Write the Anomalias sheet with index, date, value, anomaly, classification."""
    _write_header(ws, ["index", "date", "value", "anomaly", "classification"])

    row_idx = 2
    for index_name in SUPPORTED_INDICES:
        index_data = indices.get(index_name, {})
        anomalies = index_data.get("anomalies", []) or []
        for anomaly in anomalies:
            if anomaly.get("anomaly") is not True:
                continue
            date_value = _to_iso_date(anomaly.get("date"))
            value = _format_number_for_cell(anomaly.get("value"))
            classification = anomaly.get("classification")

            ws.cell(row=row_idx, column=1, value=index_name)
            ws.cell(row=row_idx, column=2, value=date_value)
            cell_valor = ws.cell(row=row_idx, column=3, value=value)
            if value is not None:
                cell_valor.number_format = "0.00000000"
            ws.cell(row=row_idx, column=4, value=True)
            ws.cell(row=row_idx, column=5, value=classification)
            row_idx += 1


def _write_geografia_sheet(ws, context: dict) -> None:
    """Write the Geografia sheet with field, value, unit."""
    _write_header(ws, ["field", "value", "unit"])

    area_ha = context.get("area_ha")
    map_center = context.get("map_center")
    geometry = context.get("geometry_geojson", {})

    rows = [
        ("area_ha", area_ha, "ha" if area_ha is not None else ""),
    ]

    if map_center and isinstance(map_center, (list, tuple)) and len(map_center) >= 2:
        rows.append(("centroid_latitude", map_center[0], "graus"))
        rows.append(("centroid_longitude", map_center[1], "graus"))

    crs = geometry.get("crs", {}).get("properties", {}).get("name", "EPSG:4326")
    rows.append(("crs", crs, "" if crs else ""))

    # GeoJSON as string
    geojson_str = json.dumps(geometry, ensure_ascii=False, separators=(",", ":"))
    rows.append(("geojson", geojson_str, ""))

    for row_idx, (field, value, unit) in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=field)
        cell_valor = ws.cell(row=row_idx, column=2, value=_format_number_for_cell(value))
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell_valor.number_format = "0.00"
        # For non-numeric values (like geojson string), write as string
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            ws.cell(row=row_idx, column=2, value=value)
        ws.cell(row=row_idx, column=3, value=unit)

    # Set column widths for readability
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 20


def render_analysis_xlsx(
    export_context: dict,
    *,
    generated_at: date | datetime | None = None,
) -> bytes:
    """Render a multi-sheet XLSX workbook from a normalized export context.

    The workbook contains exactly 7 sheets in fixed order:
    - Resumo: campo, valor, unidade
    - NDVI, NDWI, NDMI: date, value, trend, anomaly
    - Clima: date, precipitation_mm, temperature_c, solar_radiation_mj_m2_day
    - Anomalias: index, date, value, anomaly, classification
    - Geografia: field, value, unit

    Args:
        export_context: Normalized analysis snapshot from build_export_context.
        generated_at: Optional generation timestamp for the report.

    Returns:
        Bytes of the XLSX workbook.

    Raises:
        ValueError: If context is invalid or missing required indices.
    """
    if not isinstance(export_context, dict):
        raise ValueError("O contexto de exportacao deve ser um dicionario")

    indices = export_context.get("indices")
    if not isinstance(indices, dict) or any(name not in indices for name in SUPPORTED_INDICES):
        raise ValueError("O contexto deve conter NDVI, NDWI e NDMI")

    generated_at = generated_at or datetime.now()

    wb = openpyxl.Workbook()
    wb.properties.title = "Relatório de Monitoramento Agrícola"
    wb.properties.creator = "Monitoramento Agrícola por Satélite"
    wb.properties.created = generated_at

    # Remove default sheet
    wb.remove(wb.active)

    # Sheet 1: Resumo
    ws_resumo = wb.create_sheet("Resumo")
    _write_resumo_sheet(ws_resumo, export_context, generated_at)
    ws_resumo.column_dimensions["A"].width = 35
    ws_resumo.column_dimensions["B"].width = 20
    ws_resumo.column_dimensions["C"].width = 20

    # Sheets 2-4: Index sheets (NDVI, NDWI, NDMI)
    for index_name in SUPPORTED_INDICES:
        ws = wb.create_sheet(index_name)
        _write_index_sheet(ws, index_name, indices[index_name])
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 12

    # Sheet 5: Clima
    ws_clima = wb.create_sheet("Clima")
    _write_clima_sheet(ws_clima, export_context.get("climate_data"))
    ws_clima.column_dimensions["A"].width = 15
    ws_clima.column_dimensions["B"].width = 20
    ws_clima.column_dimensions["C"].width = 20
    ws_clima.column_dimensions["D"].width = 30

    # Sheet 6: Anomalias
    ws_anomalias = wb.create_sheet("Anomalias")
    _write_anomalias_sheet(ws_anomalias, indices)
    ws_anomalias.column_dimensions["A"].width = 12
    ws_anomalias.column_dimensions["B"].width = 15
    ws_anomalias.column_dimensions["C"].width = 15
    ws_anomalias.column_dimensions["D"].width = 12
    ws_anomalias.column_dimensions["E"].width = 18

    # Sheet 7: Geografia
    ws_geografia = wb.create_sheet("Geografia")
    _write_geografia_sheet(ws_geografia, export_context)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()